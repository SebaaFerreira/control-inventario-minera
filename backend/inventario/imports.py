import json
import unicodedata
from django.core import serializers
from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import transaction, IntegrityError
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError
from .models import Bodega, Categoria, Articulo, Movimiento, Trabajador
from .serializers import TrabajadorSerializer


def restaurar_respaldo(archivo):
    if archivo.size > 20 * 1024 * 1024:
        raise ValidationError('El respaldo supera 20 MB.')
    try:
        datos = json.load(archivo)
        modelos = {m._meta.label_lower: m for m in (Bodega, Categoria, Articulo, Movimiento, Trabajador)}
        if not isinstance(datos, list) or not datos:
            raise ValueError('El respaldo debe contener una lista de registros de inventario.')
        claves = set()
        for fila in datos:
            if not isinstance(fila, dict) or fila.get('model') not in modelos or not isinstance(fila.get('fields'), dict):
                raise ValueError('Solo se permiten registros de inventario.')
            if type(fila.get('pk')) is not int or fila['pk'] <= 0:
                raise ValueError('Cada registro debe tener un ID positivo.')
            clave = (fila['model'], fila['pk'])
            if clave in claves:
                raise ValueError('El respaldo contiene IDs duplicados.')
            claves.add(clave)
        for fila in datos:
            for campo in modelos[fila['model']]._meta.fields:
                if campo.is_relation:
                    valor = fila['fields'].get(campo.name)
                    if valor is not None and (campo.related_model._meta.label_lower, valor) not in claves:
                        raise ValueError(f'Referencia faltante: {campo.name}.')
        objetos = list(serializers.deserialize('json', json.dumps(datos)))
        # Guardado raw: el snapshot ya contiene el stock; no se repiten movimientos.
        with transaction.atomic():
            Movimiento.objects.all().delete()
            Articulo.objects.all().delete()
            Trabajador.objects.all().update(capataz_asignado=None, supervisor_asignado=None)
            Trabajador.objects.all().delete()
            Categoria.objects.all().delete()
            Bodega.objects.all().delete()
            for objeto in objetos:
                objeto.save()
            for objeto in objetos:
                objeto.object.full_clean()
                if isinstance(objeto.object, Articulo) and objeto.object.estado not in ('OPERATIVO', 'MANTENIMIENTO', 'BAJA'):
                    raise ValueError('Estado técnico inválido en el respaldo.')
    except (ValueError, TypeError, KeyError, IntegrityError, DjangoValidationError, serializers.base.DeserializationError) as exc:
        raise ValidationError({'error': f'Respaldo inválido. No se modificaron los datos: {exc}'}) from exc


def normalizar(texto):
    return ''.join(c for c in unicodedata.normalize('NFD', str(texto).strip().upper()) if unicodedata.category(c) != 'Mn')


def importar_trabajadores(archivo):
    if not archivo.name.lower().endswith(('.xlsx', '.xlsm')):
        raise ValidationError('Use un archivo .xlsx o .xlsm; .xls no está soportado.')
    if archivo.size > 10 * 1024 * 1024:
        raise ValidationError('La planilla supera 10 MB.')
    try:
        libro = load_workbook(archivo, data_only=True, read_only=True)
    except Exception as exc:
        raise ValidationError('No se pudo abrir la planilla Excel.') from exc
    try:
        hoja = None
        indices = {}
        for sheet in libro.worksheets:
            for numero, fila in enumerate(sheet.iter_rows(max_row=50, values_only=True), start=1):
                titulos = [normalizar(v or '') for v in fila]
                if any('RUT' in v for v in titulos) and any('NOMBRE' in v for v in titulos):
                    hoja, inicio = sheet, numero + 1
                    for i, titulo in enumerate(titulos):
                        if 'RUT' in titulo: indices['rut'] = i
                        elif 'NOMBRE' in titulo: indices['nombre_completo'] = i
                        elif 'CARGO' in titulo or 'ESPECIALIDAD' in titulo: indices['especialidad'] = i
                        elif 'ASISTENCIA' in titulo: indices['estado_asistencia'] = i
                        elif 'CICLO' in titulo: indices['turno_asignado'] = i
                        elif 'SISTEMA' in titulo or 'TURNO' in titulo: indices['sistema_turno'] = i
                        elif 'TELEFONO' in titulo: indices['telefono'] = i
                        elif 'TEST' in titulo: indices['test_esfuerzo'] = i
                        elif 'HABITACION' in titulo: indices['habitacion'] = i
                    break
            if hoja is not None:
                break
        if hoja is None:
            raise ValidationError('No se encontró una tabla con RUT y NOMBRE en las primeras 50 filas.')
        preparados = []
        vistos = set()
        for numero, fila in enumerate(hoja.iter_rows(min_row=inicio, values_only=True), start=inicio):
            valores = {k: str(fila[i]).strip() if i < len(fila) and fila[i] is not None else '' for k, i in indices.items()}
            if not valores['rut']:
                continue
            rut = valores['rut'].replace('.', '').replace(' ', '').upper()
            if rut in vistos:
                raise ValidationError(f'Fila {numero}: RUT duplicado en la planilla.')
            vistos.add(rut)
            valores['rut'] = rut
            if 'turno_asignado' in valores:
                turno = normalizar(valores['turno_asignado']).replace('TURNO ', '')
                valores['turno_asignado'] = turno or 'DIA'
            elif valores.get('sistema_turno') in ('A', 'B', 'E', 'DIA', 'NOCHE'):
                valores['turno_asignado'] = valores.pop('sistema_turno')
            if 'especialidad' in valores:
                cargo = normalizar(valores['especialidad'])
                valores['rol'] = 'CAPATAZ' if 'CAPATAZ' in cargo else 'SUPERVISOR' if 'SUPERVISOR' in cargo or 'JEFE' in cargo else 'OPERARIO'
            preparados.append((numero, valores))
        if not preparados:
            raise ValidationError('La planilla no contiene trabajadores debajo del encabezado.')
        creados = actualizados = 0
        with transaction.atomic():
            existentes = {t.rut.replace('.', '').replace(' ', '').upper(): t for t in Trabajador.objects.all()}
            for numero, valores in preparados:
                trabajador = existentes.get(valores['rut'])
                serializer = TrabajadorSerializer(trabajador, data=valores, partial=trabajador is not None)
                if not serializer.is_valid():
                    raise ValidationError({'error': f'Fila {numero}: {serializer.errors}. No se importó ninguna fila.'})
                serializer.save()
                if trabajador is None: creados += 1
                else: actualizados += 1
        return creados, actualizados
    finally:
        libro.close()
