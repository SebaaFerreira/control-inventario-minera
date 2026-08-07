import os
import tempfile
import openpyxl
from io import StringIO
from django.core.management import call_command
from django.http import HttpResponse, JsonResponse
from rest_framework.decorators import api_view
from rest_framework import viewsets
from .models import Bodega, Categoria, Articulo, Movimiento, Trabajador
from .serializers import BodegaSerializer, CategoriaSerializer, ArticuloSerializer, MovimientoSerializer, TrabajadorSerializer

class BodegaViewSet(viewsets.ModelViewSet):
    queryset = Bodega.objects.all()
    serializer_class = BodegaSerializer

class CategoriaViewSet(viewsets.ModelViewSet):
    queryset = Categoria.objects.all()
    serializer_class = CategoriaSerializer

class ArticuloViewSet(viewsets.ModelViewSet):
    queryset = Articulo.objects.all()
    serializer_class = ArticuloSerializer
    
    def get_queryset(self):
        queryset = Articulo.objects.all()
        categoria = self.request.query_params.get('categoria', None)
        if categoria is not None:
            queryset = queryset.filter(categoria__iexact=categoria)
        return queryset

class MovimientoViewSet(viewsets.ModelViewSet):
    queryset = Movimiento.objects.all().order_by('-id')
    serializer_class = MovimientoSerializer

    def perform_create(self, serializer):
        movimiento = serializer.save()
        if movimiento.tipo_movimiento == 'SALIDA':
            articulo = movimiento.articulo
            articulo.stock_actual -= movimiento.cantidad
            if articulo.stock_actual < 0:
                articulo.stock_actual = 0
            articulo.save()

    def perform_update(self, serializer):
        movimiento_viejo = self.get_object()
        estaba_devuelto = (movimiento_viejo.estado_prestamo == 'DEVUELTO')
        
        movimiento_nuevo = serializer.save()
        
        if not estaba_devuelto and movimiento_nuevo.estado_prestamo == 'DEVUELTO':
            articulo = movimiento_nuevo.articulo
            articulo.stock_actual += movimiento_nuevo.cantidad
            articulo.save()

class TrabajadorViewSet(viewsets.ModelViewSet):
    queryset = Trabajador.objects.all()
    serializer_class = TrabajadorSerializer


# =========================================================
# 📥 LÓGICA DE EXPORTACIÓN E IMPORTACIÓN (JSON)
# =========================================================
@api_view(['GET'])
def exportar_respaldo(request):
    try:
        out = StringIO()
        call_command('dumpdata', 'inventario', format='json', indent=4, stdout=out)
        response = HttpResponse(out.getvalue(), content_type='application/json')
        response['Content-Disposition'] = 'attachment; filename="respaldo_bodega.json"'
        return response
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)

@api_view(['POST'])
def importar_respaldo(request):
    try:
        if 'archivo' not in request.FILES:
            return JsonResponse({'error': 'No se adjuntó archivo.'}, status=400)
        archivo = request.FILES['archivo']
        with tempfile.NamedTemporaryFile(delete=False, suffix='.json') as tmp:
            for chunk in archivo.chunks(): tmp.write(chunk)
            tmp_path = tmp.name
        call_command('loaddata', tmp_path)
        os.remove(tmp_path)
        return JsonResponse({'mensaje': 'Base de datos restaurada.'})
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)


# =========================================================
# 🟢 MEGA-IMPORTADOR INTELIGENTE (TARJA EXCEL MARICUNGA)
# =========================================================
@api_view(['POST'])
def importar_excel_trabajadores(request):
    try:
        if 'archivo' not in request.FILES:
            return JsonResponse({'error': 'No se adjuntó ningún archivo.'}, status=400)

        archivo = request.FILES['archivo']
        # data_only=True extrae el texto final y omite las fórmulas macro
        wb = openpyxl.load_workbook(archivo, data_only=True)

        indices = {}
        creados = 0
        actualizados = 0
        hoja_datos = None
        fila_inicio = 0

        # 1. ESCANEAR TODAS LAS HOJAS DEL LIBRO (El problema era wb.active)
        for sheet in wb.worksheets:
            # Extraemos las primeras 50 filas para analizarlas rápido
            filas_prueba = list(sheet.iter_rows(values_only=True, max_row=50))
            
            for idx_fila, row in enumerate(filas_prueba):
                # Limpiamos todos los valores de la fila
                row_strs = [str(c).strip().upper() if c is not None else "" for c in row]
                
                # ¿Están las columnas mágicas en esta fila?
                if any("RUT" in c for c in row_strs) and any("NOMBRE" in c for c in row_strs):
                    hoja_datos = sheet
                    fila_inicio = idx_fila + 1 # Empezamos a leer la fila que sigue (los datos reales)
                    
                    # Mapeamos los índices exactos de la foto de tu Excel
                    for idx_col, val in enumerate(row_strs):
                        if "RUT" in val: indices['rut'] = idx_col
                        elif "NOMBRE" in val: indices['nombre'] = idx_col
                        elif "CARGO" in val: indices['cargo'] = idx_col
                        elif "ESTADO" in val and "ASISTENCIA" in val: indices['estado'] = idx_col
                        elif "TURNO" in val: indices['turno'] = idx_col
                        elif "CICLO" in val: indices['ciclo'] = idx_col
                        elif "TELEFONO" in val: indices['telefono'] = idx_col
                        elif "TEST" in val: indices['test'] = idx_col
                        elif "HABITACION" in val: indices['habitacion'] = idx_col
                    break # Rompe el ciclo de filas
            if hoja_datos:
                break # Rompe el ciclo de hojas porque ya encontramos la tabla

        # Si después de revisar todo el libro no hay cabeceras
        if not hoja_datos:
            nombres_hojas = ", ".join(wb.sheetnames)
            return JsonResponse({'error': f'No se encontró la tabla con RUT y NOMBRE en ninguna de estas pestañas: {nombres_hojas}'}, status=400)

        # 2. PROCESAR LA HOJA CORRECTA
        filas_completas = list(hoja_datos.iter_rows(values_only=True))
        
        for row in filas_completas[fila_inicio:]:
            def get_val(key):
                if key in indices and indices[key] < len(row):
                    val = row[indices[key]]
                    v_str = str(val).strip() if val is not None else ""
                    return "" if v_str.upper() == 'NONE' else v_str
                return ""

            rut_val = get_val('rut')
            # Ignoramos si está vacio
            if not rut_val or rut_val == '':
                continue

            nombre = get_val('nombre') or 'Sin Nombre'
            cargo = get_val('cargo')
            
            rol_calc = 'OPERARIO'
            if 'CAPATAZ' in cargo.upper(): rol_calc = 'CAPATAZ'
            elif 'JEFE' in cargo.upper() or 'SUPERVISOR' in cargo.upper(): rol_calc = 'SUPERVISOR'

            defaults = {
                'nombre_completo': nombre[:150],
                'rol': rol_calc,
                'especialidad': cargo[:100],
                'estado_asistencia': get_val('estado')[:50],
                'sistema_turno': get_val('turno')[:50],
                'turno_asignado': get_val('ciclo')[:20] or 'DIA',
                'telefono': get_val('telefono')[:50],
                'test_esfuerzo': get_val('test')[:50],
                'habitacion': get_val('habitacion')[:100],
                'activo': True
            }

            _, created = Trabajador.objects.update_or_create(rut=rut_val[:12], defaults=defaults)
            if created: creados += 1
            else: actualizados += 1

        if creados == 0 and actualizados == 0:
            return JsonResponse({'error': 'Encontramos los títulos, pero no pudimos leer ningún RUT válido debajo de ellos.'}, status=400)

        return JsonResponse({'mensaje': f'Trabajadores Nuevos: {creados} | Trabajadores Actualizados: {actualizados}'})

    except Exception as e:
        return JsonResponse({'error': f"Error técnico al leer el archivo: {str(e)}"}, status=500)